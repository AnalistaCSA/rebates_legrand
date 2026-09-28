
import openpyxl
from collections import defaultdict

# =============================================================
# 1. CARGAR ARCHIVO EXCEL
# =============================================================

archivo = 'Compras Legrand.xlsx'

wb = openpyxl.load_workbook(archivo, data_only=True)

facturado = wb['Facturado']
compras = wb['Compras']

encabezados_facturado = [
    celda.value for celda in facturado[1]
]

encabezados_compras = [
    celda.value for celda in compras[1]
]


# =============================================================
# 2. CONFIGURACIÓN DE TRIMESTRES
# =============================================================

meses_iniciales_trimestre = {
    'Q1': ['ENERO', 'FEBRERO'],
    'Q2': ['ABRIL', 'MAYO'],
    'Q3': ['JULIO', 'AGOSTO'],
    'Q4': ['OCTUBRE', 'NOVIEMBRE']
}


# =============================================================
# 3. FUNCIÓN PARA CALCULAR REBATES
# =============================================================

def calcular_rebates(trimestre, categoria, meta, compras_usadas):

    total_facturado = 0
    total_compras = 0
    nuevas_compras_usadas = []

    # =========================================================
    # 4. CALCULAR FACTURACIÓN REAL DEL TRIMESTRE
    # =========================================================

    for fila in facturado.iter_rows(
        min_row=2,
        values_only=True
    ):

        trimestre_facturado = fila[
            encabezados_facturado.index('Trimestre')
        ]

        categoria_facturado = fila[
            encabezados_facturado.index('Categoria')
        ]

        subcategoria_facturada = (
            fila[encabezados_facturado.index('Subcategoria')]
            if 'Subcategoria' in encabezados_facturado
            else None
        )

        subtotal_facturado = (
            fila[
                encabezados_facturado.index('Subtotal')
            ] or 0
        )

        # Electrobarras se identifica por subcategoría.
        if categoria == 'ELECTROBARRAS':

            cumple_categoria = (
                subcategoria_facturada == 'ELECTROBARRAS'
            )

        else:

            cumple_categoria = (
                categoria_facturado == categoria
            )

        if (
            trimestre_facturado == trimestre
            and cumple_categoria
        ):

            total_facturado += float(subtotal_facturado)

    # =========================================================
    # 5. CALCULAR EL BACK ORDER RECONOCIDO
    # =========================================================
    # El Back Order elegible se debe reconocer aunque la facturación
    # real ya haya alcanzado la meta, porque también hace parte de la
    # base sobre la que se calcula el rebate.

    # =========================================================
    # 6. AGRUPAR COMPRAS ELEGIBLES POR OC
    # =========================================================

    compras_por_oc = defaultdict(float)
    filas_por_oc = defaultdict(list)

    for numero_fila, compra in enumerate(
        compras.iter_rows(
            min_row=2,
            values_only=True
        ),
        start=2
    ):

        trimestre_compra = compra[
            encabezados_compras.index('Trimestre')
        ]

        categoria_compra = compra[
            encabezados_compras.index('Categoria')
        ]

        mes_compra = compra[
            encabezados_compras.index('Mes')
        ]

        subtotal_compra = (
            compra[
                encabezados_compras.index('Subtotal Real')
            ] or 0
        )

        oc_compra = compra[
            encabezados_compras.index('OC')
        ]

        subcategoria_compra = (
            compra[encabezados_compras.index('Subcategoria')]
            if 'Subcategoria' in encabezados_compras
            else None
        )

        # -----------------------------------------------------
        # Validar categoría.
        # Electrobarras se filtra por subcategoría.
        # -----------------------------------------------------

        if categoria == 'ELECTROBARRAS':

            cumple_categoria = (
                subcategoria_compra == 'ELECTROBARRAS'
            )

        else:

            cumple_categoria = (
                categoria_compra == categoria
            )

        # -----------------------------------------------------
        # Validar trimestre y meses elegibles.
        # -----------------------------------------------------

        if (
            trimestre_compra != trimestre
            or not cumple_categoria
            or mes_compra not in meses_iniciales_trimestre[trimestre]
            or not oc_compra
        ):
            continue

        # -----------------------------------------------------
        # Agrupar todas las líneas de una misma OC.
        # -----------------------------------------------------

        compras_por_oc[oc_compra] += float(subtotal_compra)

        filas_por_oc[oc_compra].append({
            'fila': numero_fila,
            'subtotal': float(subtotal_compra)
        })

    # =========================================================
    # 7. CALCULAR BACK ORDER DE CADA OC
    # =========================================================

    for oc_compra, importe_compra in compras_por_oc.items():

        # -----------------------------------------------------
        # Calcular facturación asociada a la OC.
        # -----------------------------------------------------

        facturado_de_esta_compra = 0

        for fila_facturado in facturado.iter_rows(
            min_row=2,
            values_only=True
        ):

            oc_facturada = fila_facturado[
                encabezados_facturado.index('OC')
            ]

            categoria_facturada = fila_facturado[
                encabezados_facturado.index('Categoria')
            ]

            trimestre_compra_facturado = fila_facturado[
                encabezados_facturado.index('Trimestre Compra')
            ]

            subtotal_facturado = (
                fila_facturado[
                    encabezados_facturado.index('Subtotal')
                ] or 0
            )

            subcategoria_facturada = (
                fila_facturado[
                    encabezados_facturado.index('Subcategoria')
                ]
                if 'Subcategoria' in encabezados_facturado
                else None
            )

            # -------------------------------------------------
            # Validar categoría de la factura.
            # -------------------------------------------------

            if categoria == 'ELECTROBARRAS':

                cumple_categoria_facturada = (
                    subcategoria_facturada == 'ELECTROBARRAS'
                )

            else:

                cumple_categoria_facturada = (
                    categoria_facturada == categoria
                )

            if (
                oc_facturada == oc_compra
                and cumple_categoria_facturada
                and trimestre_compra_facturado == trimestre
            ):

                facturado_de_esta_compra += float(
                    subtotal_facturado
                )

        # =====================================================
        # 8. CALCULAR BACK ORDER DISPONIBLE
        # =====================================================

        back_order = max(
            0,
            importe_compra - facturado_de_esta_compra
        )

        if back_order <= 0:
            continue

        # =====================================================
        # 9. DESCONTAR VALORES UTILIZADOS ANTERIORMENTE
        # =====================================================

        valor_ya_utilizado = sum(
            compras_usadas.get(
                fila_compra['fila'],
                0
            )
            for fila_compra in filas_por_oc[oc_compra]
        )

        back_order_disponible = max(
            0,
            back_order - valor_ya_utilizado
        )

        if back_order_disponible <= 0:
            continue

        # =====================================================
        # 10. SUMAR TODO EL BACK ORDER DISPONIBLE
        # =====================================================

        valor_a_utilizar = back_order_disponible

        total_compras += valor_a_utilizar

        # =====================================================
        # 11. REGISTRAR EL VALOR UTILIZADO POR FILA
        # =====================================================

        saldo_por_distribuir = valor_a_utilizar

        for fila_compra in filas_por_oc[oc_compra]:

            numero_fila = fila_compra['fila']

            valor_anterior = compras_usadas.get(
                numero_fila,
                0
            )

            saldo_fila = max(
                0,
                fila_compra['subtotal'] - valor_anterior
            )

            valor_fila_utilizado = min(
                saldo_fila,
                saldo_por_distribuir
            )

            if valor_fila_utilizado <= 0:
                continue

            compras_usadas[numero_fila] = (
                valor_anterior + valor_fila_utilizado
            )

            saldo_por_distribuir -= valor_fila_utilizado

            nuevas_compras_usadas.append({
                'fila': numero_fila,
                'oc': oc_compra,
                'valor_utilizado': valor_fila_utilizado,
                'back_order_original': back_order,
                'back_order_restante': max(
                    0,
                    back_order - valor_a_utilizar
                )
            })

            if saldo_por_distribuir <= 0:
                break

    # =========================================================
    # 12. CALCULAR BASE PARA CUMPLIMIENTO Y REBATE
    # =========================================================

    total = total_facturado + total_compras

    # =========================================================
    # 13. DEVOLVER RESULTADOS
    # =========================================================

    return (
        total_facturado,
        total_compras,
        total,
        nuevas_compras_usadas
    )